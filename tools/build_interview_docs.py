from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "interview-prep"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1C2526"
MUTED = "5F6B6D"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "116247"
GOLD = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(element, fill: str):
    props = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run("SMART CAMERA SURVEILLANCE  |  INTERVIEW PREPARATION")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D7DBE2")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    add_page_number(section.footer.paragraphs[0])
    doc.core_properties.author = "Smart Camera Surveillance Project"
    doc.core_properties.subject = "Project-specific interview preparation"
    doc.core_properties.keywords = "surveillance, WebRTC, RTSP, YOLO, Redis Streams, Postgres, Docker"


def add_cover(doc: Document, title: str, subtitle: str, guide_number: str, note: str):
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run(f"INTERVIEW PREPARATION  /  {guide_number}")
    set_run_font(run, size=10, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(title)
    set_run_font(run, size=28, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run(subtitle)
    set_run_font(run, size=14, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Real-Time Camera Surveillance Dashboard")
    set_run_font(run, size=12, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(46)
    run = p.add_run("React + TypeScript | Bun + Hono | Python + YOLOv8n | Postgres | Redis | MediaMTX | Docker")
    set_run_font(run, size=9.5, color=MUTED)

    add_callout(doc, "Implementation note", note, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("Prepared from the implemented codebase | September 2026")
    set_run_font(run, size=9, color=MUTED, italic=True)
    doc.add_page_break()


def add_callout(doc: Document, label: str, text: str, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    shade(p._p, CALLOUT)
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True, color=color)
    r = p.add_run(text)
    set_run_font(r, color=INK)


def add_body(doc: Document, text: str, bold_start: str | None = None):
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        r = p.add_run(bold_start)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_start):])
        set_run_font(r)
    else:
        set_run_font(p.add_run(text))
    return p


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item))


def add_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item))


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    shade(p._p, LIGHT_GRAY)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        set_run_font(p.add_run(line), name="Consolas", size=8.8, color=INK)
    return p


def add_labeled(doc: Document, label: str, text: str, color=DARK_BLUE):
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True, color=color)
    set_run_font(p.add_run(text))
    return p


def add_question(doc: Document, number: int, difficulty: str, question: str, answer: str, anchors: str):
    doc.add_heading(f"Q{number}. {question}", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    label_color = GREEN if difficulty == "Basic" else GOLD if difficulty == "Intermediate" else RED
    r = p.add_run(difficulty.upper())
    set_run_font(r, size=8.5, bold=True, color=label_color)
    r = p.add_run(f"  |  Code anchors: {anchors}")
    set_run_font(r, size=8.5, color=MUTED, italic=True)
    add_labeled(doc, "Answer", answer)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        shade(cell._tc, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=9.5, bold=True, color=DARK_BLUE)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=9.2)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def save(doc: Document, filename: str):
    path = OUT / filename
    doc.save(path)
    return path


def build_questions():
    doc = Document()
    configure_styles(doc)
    add_cover(
        doc,
        "Interview Questions & Answers",
        "37 answers you can say out loud, grounded in the actual code",
        "GUIDE 01",
        "The implemented worker is Python, not Go. These answers describe Python, OpenCV, FFmpeg, and Ultralytics exactly as used in the repository.",
    )
    doc.add_heading("How To Use This Guide", level=1)
    add_bullets(doc, [
        "Read the answer once, close the document, and repeat the idea in your own words.",
        "Practice the basic questions first, then add the intermediate and deep-dive details.",
        "Do not claim model training, cloud deployment, snapshots, alert-history UI, or multi-worker scaling; those are not implemented.",
        "Use the code anchors to connect each answer to a real file before an interview.",
    ])
    add_callout(doc, "Best opening", "Start with the user journey, then explain why the system is split into an API, video server, worker, database, message stream, and browser UI.")

    sections = [
        ("Project Overview & Architecture", [
            ("Basic", "Walk me through your project.", "I built a local real-time camera surveillance dashboard. A user can sign up or log in, register an RTSP camera, start or stop processing, watch the stream in a browser, and see person-detection alerts and live statistics. The browser is React and TypeScript. A Bun and Hono API handles authentication, camera data, alerts, and WebSocket fan-out. Postgres stores users, cameras, alerts, and latest camera stats. Redis Streams carries start and stop commands to a Python worker and carries detections, stats, and state back. The worker uses OpenCV and a pretrained YOLOv8n model, and MediaMTX bridges RTSP video to browser-compatible WebRTC. Docker Compose starts the seven local services together.", "README.md; docker-compose.yml"),
            ("Intermediate", "What problem does the architecture solve?", "The main problem is that camera processing is long-running and browsers cannot directly play normal RTSP security-camera streams. I separated short web requests from video processing. The API accepts a start request quickly and publishes a Redis command instead of blocking while it reads frames. The Python worker owns the long-running camera loop and ML inference. MediaMTX handles the protocol boundary between RTSP and WebRTC. Postgres stores durable business data, while Redis carries transient commands and events. That separation keeps the UI responsive and makes failures easier to isolate by layer.", "api/src/routes/camera-routes.ts; worker/worker.py; frontend/src/lib/whep.ts"),
            ("Deep dive", "What happens step by step when I start a camera?", "First, the React dashboard sends POST /cameras/:id/start with a bearer token. The API verifies the JWT, updates that user's camera to desired state live and runtime state connecting, then writes a start command to camera.commands in Redis. The Python worker consumes and acknowledges the command, creates a CameraRunner, starts an FFmpeg restream to a camera-specific MediaMTX path, and opens the source RTSP URL with OpenCV. Once the input opens, it publishes a live state and periodic stats. The browser requests the camera's WHEP URL from MediaMTX and receives WebRTC video. In parallel, YOLO processes selected frames and the worker publishes person events for the API to deduplicate, store, and broadcast.", "api/src/routes/camera-routes.ts; worker/worker.py; worker/camera_runner.py"),
            ("Intermediate", "Why did you use separate services instead of one application?", "Each service has a different responsibility and runtime profile. The API handles short request-response work, the worker performs CPU-heavy video and inference work, MediaMTX specializes in video protocols, and Postgres and Redis provide storage and messaging. Keeping those concerns separate means a camera failure does not have to crash authentication, and I can inspect or restart one layer independently. The cost is operational complexity: there are more containers, ports, environment variables, and failure paths. For this project, Compose keeps that complexity manageable while still demonstrating a production-shaped design.", "docker-compose.yml; docs/architecture.md"),
        ]),
        ("Frontend: React, TypeScript & Realtime UI", [
            ("Basic", "What does the frontend do?", "The frontend provides authentication, camera creation and editing, start and stop controls, live video, recent alerts, FPS, detections per minute, and stream diagnostics. App.tsx mainly composes the page and passes handlers. useDashboard owns session storage, REST loading, WebSocket events, and camera actions. CameraTile presents one camera, and WebRtcVideo owns WHEP negotiation and retry behavior. I used TypeScript types for cameras, alerts, runtime states, and WebSocket message variants so data mismatches are caught earlier.", "frontend/src/App.tsx; frontend/src/hooks/useDashboard.ts; frontend/src/types.ts"),
            ("Intermediate", "Why did you move dashboard logic into a custom hook?", "The first version could easily become one large component mixing UI markup, authentication state, REST calls, WebSocket lifecycle, and updates. I moved those responsibilities into useDashboard so App stays focused on rendering and composition. The hook gives one place for session persistence, camera and alert collections, refresh behavior, start and stop actions, and realtime event handling. This is not abstraction for its own sake; it separates stateful behavior from presentation and makes both easier to test and change.", "frontend/src/hooks/useDashboard.ts; frontend/src/App.tsx"),
            ("Intermediate", "Why are REST, WebSocket, and WebRTC all used?", "They solve three different communication problems. REST is best for explicit commands and queries such as login, list cameras, create a camera, and start processing. WebSocket is for small server-pushed events such as new alerts, FPS, and runtime state, so the browser does not poll repeatedly. WebRTC is optimized for low-latency audio and video delivery. I would not send video frames over the application's WebSocket because that would duplicate work MediaMTX already handles much better.", "frontend/src/lib/api.ts; frontend/src/hooks/useDashboard.ts; frontend/src/lib/whep.ts"),
            ("Deep dive", "How does WebSocket reconnection work?", "The dashboard creates a token-authenticated WebSocket. If it closes, the hook schedules another connection with exponential backoff: one second, then two, four, eight, and a maximum delay of ten seconds. The top bar shows connecting, reconnecting, connected, or disconnected. When the API confirms the connection, I reset the retry count and call refresh so camera and alert data missed during the disconnection is loaded from REST. Cleanup closes the socket and clears the timer when the user logs out or the component unmounts, which prevents duplicate connections.", "frontend/src/hooks/useDashboard.ts"),
            ("Deep dive", "What do the stream diagnostics tell you?", "The tile deliberately separates requested, worker, and browser state. Requested is the desired state stored by the API. Worker is the runtime state reported by the Python process. Browser is measured by the WebRTC component and can be connecting, retrying, or playing. That separation helped diagnose a real issue where MediaMTX and the worker were healthy but the UI state was stale after an API restart. FPS, input enabled state, and the latest error give additional evidence without requiring the user to open every container log.", "frontend/src/components/CameraTile.tsx; frontend/src/components/WebRtcVideo.tsx"),
        ]),
        ("API, Authentication & Ownership", [
            ("Basic", "What does Bun and Hono do in this project?", "Bun is the JavaScript and TypeScript runtime used to start the backend and run its tests. Hono is the small web framework that defines middleware and routes. Hono handles paths such as /health, /auth, /cameras, /alerts, and /ws, while Bun supplies the server and WebSocket integration. I kept route handlers grouped by feature and moved SQL into repository modules so HTTP concerns and database concerns are not mixed everywhere.", "api/src/index.ts; api/src/routes; api/src/repositories"),
            ("Intermediate", "How is signup and login implemented?", "Signup validates the username and password, hashes the password with Argon2, inserts the user, and returns a signed JWT. Login finds the user by username, verifies the password against the stored hash, and returns a new JWT if it matches. The token contains the user ID as the subject and the username as a claim, and it expires after twelve hours. Passwords are never stored in plain text. The current version does not implement refresh tokens, account lockout, or production password policy, so I describe those as future hardening.", "api/src/routes/auth-routes.ts; api/src/auth.ts; api/src/repositories/user-repository.ts"),
            ("Intermediate", "How are protected routes secured?", "The auth middleware reads Authorization: Bearer <token>. If the token is missing it returns 401. If verification fails or the token is expired, it returns 401. On success it stores the authenticated user in the Hono context, and camera and alert handlers read that identity rather than trusting a user ID supplied by the browser. The /cameras and /alerts routes use this middleware; /health and signup/login remain public.", "api/src/auth.ts; api/src/index.ts"),
            ("Deep dive", "How do you prevent one user from controlling another user's camera?", "Authentication alone only tells me who called the API, so every camera query also enforces ownership. The repository uses both camera ID and authenticated user ID in find, update, delete, start, stop, and health queries. Lists filter by user ID, and new cameras are inserted with that owner. If the pair does not match, the API returns Camera not found rather than revealing that another user's camera exists. I added regression tests that inspect these query boundaries and HTTP tests for unowned edit, delete, start, and health behavior.", "api/src/repositories/camera-repository.ts; api/test/camera-ownership.test.ts; api/test/camera-routes.test.ts"),
            ("Intermediate", "How do you validate camera input?", "I validate in two places for different reasons. The React form gives immediate field-level feedback, but client validation can be bypassed. The API therefore uses a Zod schema as the real trust boundary. It trims strings, requires a camera name, requires a valid URL, limits name and location length, and supplies defaults for location and enabled. A formatter converts raw schema issues into a stable error object with a general message and per-field messages. Both create and partial update routes reuse the schema.", "frontend/src/components/CameraForm.tsx; api/src/validation.ts; api/src/routes/camera-routes.ts"),
            ("Intermediate", "What does the camera health endpoint return?", "GET /cameras/:id/health is ownership-protected. It joins the camera with its latest stats and returns desired state, runtime state, FPS, detections per minute, last error, stats timestamp, and whether the stats are fresh. Stats are considered fresh when they are under fifteen seconds old. The endpoint derives stopped, starting, healthy, or unhealthy. It is a compact operational view; it is not a deep health check of every external network dependency.", "api/src/repositories/camera-repository.ts; api/src/routes/camera-routes.ts"),
        ]),
        ("Postgres & Alert Persistence", [
            ("Basic", "What data is stored in Postgres?", "Postgres stores four main tables. users stores usernames and password hashes. cameras stores ownership, RTSP URL, location, enabled state, a unique stream key, desired and runtime state, and errors. alerts stores person events with camera and user IDs, timestamp, confidence, bounding box, optional snapshot URL, and dedupe key. camera_stats stores the latest FPS, detection count, state, and update time for each camera. The API creates these tables on startup using idempotent CREATE TABLE IF NOT EXISTS statements.", "api/src/db.ts"),
            ("Intermediate", "Why do you store desired state and runtime state separately?", "Desired state represents what the user asked for, while runtime state represents what the worker reports. After Start, desired can be live while runtime is still connecting or error. That difference is useful for diagnostics and recovery. If I stored only one state, a temporary failure could overwrite the user's intent, or the UI could claim a camera is live just because Start was clicked. Fresh worker stats also repair runtime state to live when a state event was missed.", "api/src/db.ts; api/src/repositories/camera-repository.ts; api/src/stream-consumer.ts"),
            ("Deep dive", "How are alert queries and pagination implemented?", "The alert repository always begins with user_id as the first filter, then optionally adds camera ID, from time, and to time. It uses cursor pagination rather than offset pagination. The cursor combines occurred_at and alert ID, and the query orders by both fields descending. It requests limit plus one row to determine whether another page exists, then returns a next cursor based on the last visible row. The API supports the history data, but the current dashboard only displays the newest four alerts per camera; a full history screen is still future UI work.", "api/src/repositories/alert-repository.ts; api/src/routes/alert-routes.ts; frontend/src/components/CameraTile.tsx"),
            ("Intermediate", "How do you avoid storing the same person every frame?", "The worker may detect the same person repeatedly, especially in the looped demo image. The API checks for a recent alert from the same camera and event type inside a configurable time window, currently ten seconds. Only when no recent match exists does it insert and broadcast the event. The worker also sends a time-bucket dedupe key, and event IDs are unique. This is a simple temporal dedupe strategy; production logic could track objects across frames or enforce a database uniqueness rule.", "api/src/stream-consumer.ts; worker/camera_runner.py; api/src/config.ts"),
            ("Intermediate", "Why are there indexes on alerts?", "Alert history is primarily read by user and by camera in reverse chronological order. The schema creates indexes on camera_id with occurred_at and id descending, and on user_id with the same ordering. Those columns match the filtering and cursor order used by the repository. Without the indexes, Postgres would increasingly scan and sort alert rows as history grows. I would confirm index usage with EXPLAIN ANALYZE once the table had production-scale data.", "api/src/db.ts; api/src/repositories/alert-repository.ts"),
        ]),
        ("Redis Streams & Event Flow", [
            ("Basic", "What is Redis used for?", "Redis is the message layer between the API and worker. The API writes camera.commands for start and stop. The worker writes detection.events, camera.stats, and camera.states. The API consumes those worker streams, updates Postgres, and broadcasts relevant messages to browser WebSocket clients. Redis is not the source of truth for users and alerts; Postgres is. This keeps long-running camera work out of the HTTP request lifecycle.", "api/src/config.ts; worker/config.py; docs/event-format.md"),
            ("Intermediate", "Why Redis Streams rather than Redis Pub/Sub?", "Streams retain entries, support IDs, and let the worker use a consumer group with acknowledgements. Plain Pub/Sub would drop a start command if the worker were disconnected at that moment. The worker reads new commands with XREADGROUP and acknowledges a message after handling it. On restart, it also scans recent command history and restores cameras whose latest command is start. The API's event consumer is simpler and reads new events without a consumer group, which is acceptable for this one-API local demo but would need stronger delivery semantics in a scaled deployment.", "worker/worker.py; api/src/stream-consumer.ts"),
            ("Deep dive", "How do TypeScript and Python agree on event formats?", "I documented one canonical contract and implemented matching codecs. Redis fields are string-like, so plain strings remain strings while numbers, booleans, objects, and null are JSON encoded. The API parse helper attempts JSON parsing, and the Python event helper JSON encodes non-string values. Messages consistently use camelCase fields such as cameraId, userId, occurredAt, and detectionsPerMinute. Contract tests verify representative payloads. In a larger system I would use a formal schema registry or generated types because documentation alone can drift.", "api/src/stream-codec.ts; worker/events.py; docs/event-format.md"),
            ("Deep dive", "What are the current Redis delivery limitations?", "The worker command path uses a consumer group and acknowledgements, but failed commands remain pending and there is no automatic claim strategy for a different worker. Startup recovery scans only the most recent one thousand commands and assumes one worker should restore every active camera. The API reads worker event streams from the latest position when it starts, so it can miss events produced while it was down, although persistent alerts and continuing stats reduce visible impact. For production I would add durable consumer groups for API consumers, dead-letter handling, stream trimming, idempotency constraints, and camera leases.", "worker/worker.py; api/src/stream-consumer.ts"),
        ]),
        ("Python Worker & YOLOv8n", [
            ("Basic", "Is the worker written in Go?", "No. The implemented worker is Python. I chose Python because Ultralytics YOLO and OpenCV have direct, mature APIs there, so I could focus on the video and event pipeline. The worker files are main.py, worker.py, camera_runner.py, commands.py, events.py, and config.py. Go is only mentioned as a possible future optimization; I would not claim a Go worker in an interview or on my resume.", "worker/*.py; docs/decisions.md"),
            ("Intermediate", "How is worker concurrency structured?", "The worker keeps a dictionary from camera ID to a CameraRunner and an asyncio task. A start command first stops any existing runner for that camera, then creates a new task. A stop command sets the runner's stop event and waits for shutdown, with cancellation and FFmpeg cleanup as fallbacks. This gives logical isolation per camera, but the OpenCV reads and YOLO prediction are synchronous and the model is shared. It is not true CPU-parallel inference; scaling many cameras would require worker processes, camera partitioning, or an inference service.", "worker/worker.py; worker/camera_runner.py"),
            ("Deep dive", "How is YOLOv8n integrated?", "The worker loads the pretrained yolov8n.pt model through Ultralytics. It opens RTSP frames with OpenCV and runs model.predict only when the configured detection interval has elapsed, currently every 500 milliseconds. classes=[0] restricts COCO detections to person, and the default confidence threshold is 0.35. For every box, I extract confidence and x/y/width/height, create a UUID event with model and source metadata, and publish it to detection.events. I integrated inference; I did not design or train YOLO.", "worker/camera_runner.py; worker/config.py"),
            ("Basic", "Why does the demo confidence stay around 94 percent?", "The bundled fake camera loops one still image containing one person. The pixels and composition are almost identical every time YOLO runs, so the model naturally returns nearly the same confidence. That does not mean confidence is hard-coded. If I replace the image with a moving MP4 or a real camera, confidence changes as pose, size, lighting, and occlusion change. The demo is intentionally deterministic so the full pipeline can be verified reliably.", "demo-assets/person-demo-frame.png; docker-compose.yml"),
            ("Intermediate", "How are FPS and detections per minute calculated?", "The runner keeps recent frame timestamps in a deque capped at 120 entries. FPS is the number of intervals divided by elapsed time between the oldest and newest retained timestamp. Detection timestamps are kept for a rolling sixty-second window, and their count becomes detections per minute. The worker publishes stats approximately once per second. One detail is that this is detection-box count, not a unique-person tracker, so the same person can contribute repeatedly.", "worker/camera_runner.py"),
            ("Deep dive", "How does the worker recover from failures?", "Each camera runner publishes connecting, starts its FFmpeg restream, and opens the input. If opening or reading fails, it publishes error, logs the reason, stops FFmpeg, waits three seconds, and retries until stopped. If frames stop for more than ten seconds, it raises a retryable error. On worker startup, it reads recent camera command history from newest to oldest, keeps only the latest command per camera, and restores cameras whose latest action is start. This recovery is designed for one worker; distributed ownership would need leases.", "worker/camera_runner.py; worker/worker.py; worker/commands.py"),
        ]),
        ("RTSP, WebRTC & MediaMTX", [
            ("Basic", "Why can the browser not play the RTSP URL directly?", "RTSP is common for IP cameras, but normal browser video elements do not directly support it. Browsers do support WebRTC for low-latency media. MediaMTX accepts RTSP publication and exposes a WHEP endpoint that the browser can negotiate with WebRTC. This lets camera-facing components use RTSP while the browser uses a protocol it understands.", "docker-compose.yml; frontend/src/lib/whep.ts"),
            ("Deep dive", "Explain the complete video path.", "The demo FFmpeg container loops the person image, encodes H.264 Constrained Baseline at 1280 by 720 and 25 FPS, and publishes to rtsp://mediamtx:8554/testcam. When a camera starts, the worker launches another FFmpeg process that reads the registered RTSP URL over TCP, drops audio, copies H.264 without re-encoding, and publishes to a unique cam-UUID path. The API returns a WHEP URL based on that stream key. The browser creates an RTCPeerConnection, sends an SDP offer to the WHEP URL, applies the SDP answer, receives tracks, and attaches them to the video element.", "docker-compose.yml; worker/camera_runner.py; api/src/db.ts; frontend/src/lib/whep.ts"),
            ("Intermediate", "Why did WebRTC briefly fail after Stop and Start?", "Starting is an asynchronous chain. The API can return connecting before the worker has opened the source and before FFmpeg has published the camera-specific MediaMTX path. If the browser negotiates during that window, MediaMTX correctly reports that nobody is publishing yet. I handled this by retrying WHEP every two seconds, showing a friendly Connecting video overlay, and delaying the raw error for ten seconds. The diagnostics show whether the delay is in worker state or browser playback.", "frontend/src/components/WebRtcVideo.tsx; frontend/src/lib/whep.ts"),
        ]),
        ("Docker, Testing & Production Readiness", [
            ("Basic", "What does Docker Compose run?", "Compose runs Postgres, Redis, MediaMTX, the FFmpeg demo camera, the Bun API, the Python worker, and the Nginx-served React frontend. It creates a shared network, so containers use service names such as postgres, redis, and mediamtx instead of localhost. It also maps host ports, mounts the Postgres volume and demo assets, sets environment defaults, waits for Postgres and Redis health checks where configured, and restarts the worker and demo camera unless they are explicitly stopped.", "docker-compose.yml"),
            ("Intermediate", "How did you test the system?", "I used layered tests and manual acceptance checks. Bun tests cover Redis field encoding, camera validation, ownership-scoped queries, authentication behavior, camera HTTP routes, health access, and command publishing with mocked external boundaries. Vitest and Testing Library cover logged-out and logged-in dashboard smoke behavior. Pytest covers worker event encoding and latest-command recovery. Manual checks verify the parts unit tests cannot prove alone: real RTSP publication, WebRTC playback, YOLO inference, Postgres persistence, and realtime dashboard updates.", "api/test; frontend/src/App.test.tsx; worker/tests; docs/demo-script.md"),
            ("Deep dive", "What would need to change for multiple workers?", "The current worker restoration assumes one worker owns all active cameras. For multiple workers, I would introduce camera leases or deterministic partitioning so only one worker processes a camera at a time. A lease would include worker ID and expiry and would be renewed while healthy. Redis consumer names must be unique, pending commands need claiming, and model resources need capacity-aware scheduling. I would also make event consumers durable and enforce idempotency in Postgres. None of that is implemented today, so the honest description is a single-worker local architecture with a clear scale-out path.", "worker/worker.py; docker-compose.yml"),
            ("Deep dive", "Is this production-ready?", "It is production-shaped but not production-deployed. The core local workflow works and includes ownership checks, persistence, retries, diagnostics, and tests. However, defaults include a development JWT secret, permissive CORS, localhost WebRTC settings, one worker, no TURN server, no managed backups, no snapshot object storage, and no monitoring or alert delivery outside the dashboard. The alert-history API exists, but the full history UI is not built. I would present it as a strong portfolio MVP and clearly separate implemented behavior from the production roadmap.", "README.md; docs/progress.md"),
        ]),
    ]

    number = 1
    for section, questions in sections:
        doc.add_heading(section, level=1)
        for difficulty, question, answer, anchors in questions:
            add_question(doc, number, difficulty, question, answer, anchors)
            number += 1

    doc.add_heading("Last-Minute Answer Checklist", level=1)
    add_bullets(doc, [
        "Say Python worker, not Go worker.",
        "Say integrated pretrained YOLOv8n, not trained or invented YOLO.",
        "Separate REST commands, WebSocket events, and WebRTC video.",
        "State that Postgres is durable storage and Redis Streams is messaging.",
        "Call it a local portfolio MVP, not an internet-facing production deployment.",
        "Mention the missing full alert-history UI and snapshot storage honestly.",
    ])
    assert number - 1 == 37
    return save(doc, "01_Interview_Questions_and_Answers.docx")


def build_tradeoffs():
    doc = Document()
    configure_styles(doc)
    add_cover(
        doc,
        "Challenges, Trade-offs & Why This Stack",
        "STAR stories and defensible technology decisions",
        "GUIDE 02",
        "The implemented detection worker is Python. A Go worker is discussed only as an alternative, never as completed work.",
    )
    doc.add_heading("How To Answer Trade-off Questions", level=1)
    add_body(doc, "A strong answer names the constraint, explains the decision, acknowledges the cost, and says what would trigger a different choice. Avoid pretending one technology is universally better.")
    add_callout(doc, "Interview pattern", "I chose X because of this project's constraint. Y was a realistic alternative, but its cost was ____. If the scale or team changed, I would reconsider.")

    challenges = [
        ("Bridging RTSP cameras to a browser", "The source cameras use RTSP, but the React browser cannot play RTSP directly.", "Deliver low-latency browser video without implementing a media server from scratch.", "I introduced MediaMTX as the media layer. The demo publishes H.264 RTSP into MediaMTX, the worker republishes each active camera under a unique stream key, and the browser negotiates WHEP/WebRTC. I exposed the HTTP and UDP ports required by WebRTC and used a browser RTCPeerConnection rather than pushing frames through the API.", "The browser displayed the live stream at roughly 25 FPS, while the API remained focused on control and metadata rather than video bytes."),
        ("Debugging a Docker DNS and startup failure", "After containers were recreated, the dashboard showed Requested Live, Worker Connecting, Browser Retrying, and Failed to fetch.", "Find the first failing layer instead of changing unrelated API or ML code.", "I checked API health, then logs in pipeline order. FFmpeg reported that it could not resolve mediamtx. I verified container DNS, recreated the Compose network without deleting the Postgres volume, checked the source with ffprobe from inside the worker, and confirmed MediaMTX reading and publishing sessions.", "The video path recovered. The investigation also produced a reusable diagnostics panel and a pipeline-first troubleshooting guide."),
        ("Handling asynchronous stream startup", "The API could accept Start before the worker restream existed, causing short WHEP negotiation failures.", "Keep transient startup timing from looking like a permanent user-facing error.", "I represented desired and runtime state separately, returned connecting immediately, retried WHEP every two seconds, showed a Connecting video overlay, and delayed raw WebRTC errors for ten seconds. I added separate worker and browser status to the tile.", "Normal start/stop transitions became understandable, and persistent failures still surfaced with the responsible layer."),
        ("Recovering realtime state after an API restart", "Video could continue through MediaMTX while the API restart closed the browser WebSocket, leaving FPS and state stale.", "Recover event delivery automatically and show the user whether live updates are connected.", "I added capped exponential-backoff reconnection in the dashboard. On the connected message, the frontend resets its retry counter and refreshes cameras and alerts through REST to recover missed state. Cleanup prevents duplicate sockets. A top-bar indicator shows connection status.", "The dashboard no longer requires a manual refresh after a normal API restart and makes control-plane health visible."),
        ("Preventing alert floods", "A stationary person can be detected every inference interval, especially in the repeated demo frame.", "Preserve meaningful alerts without inserting and broadcasting nearly identical events every frame.", "The worker publishes detailed detections, but the API checks for the same camera and event type inside a configurable ten-second window before storing. Alerts use UUIDs, timestamps, bounding boxes, and dedupe metadata.", "Postgres and the UI receive a manageable alert cadence while detections-per-minute remains available as a separate live statistic."),
        ("Enforcing camera ownership", "A valid JWT proves identity, but a malicious user could still try another camera UUID.", "Ensure every read and mutation is tenant-scoped without revealing another user's resources.", "I passed the authenticated user from middleware into repositories and included user_id in list, find, update, delete, state-change, and health queries. Missing ownership returns the same 404 as a nonexistent camera. I added query-boundary and HTTP route tests.", "Camera operations are consistently user-scoped, and regression tests fail if the ownership condition is removed."),
        ("Recovering active cameras after a worker restart", "Redis consumer groups do not replay already acknowledged Start commands to a new worker process.", "Restore user intent after a worker restart without coupling the worker directly to Postgres.", "For the current single-worker design, startup scans the newest one thousand camera commands, keeps the latest action per camera, and restores cameras whose latest action is Start. Camera runners also retry failed input every three seconds and declare a no-frame failure after ten seconds.", "Restarting the worker can recover active cameras and temporary RTSP outages without another manual Start command. I document that multi-worker leasing is still required for scale."),
        ("Creating a deterministic end-to-end demo", "A color-bar source verified video but could not prove person detection, while public streams are unreliable for interviews.", "Create a local, repeatable source that exercises video, inference, events, persistence, and UI updates.", "I added a local person image, mounted it read-only into the FFmpeg container, encoded it as a 25 FPS H.264 RTSP stream, and used the same URL a real camera registration would use.", "The project demonstrates the full pipeline offline. I also explain that repeated pixels are why confidence stays near the same value."),
    ]

    doc.add_heading("Part I: Real Challenges in STAR Format", level=1)
    for idx, (title, situation, task, action, result) in enumerate(challenges, 1):
        doc.add_heading(f"Challenge {idx}: {title}", level=2)
        add_labeled(doc, "Situation", situation)
        add_labeled(doc, "Task", task)
        add_labeled(doc, "Action", action)
        add_labeled(doc, "Result", result, color=GREEN)

    choices = [
        ("Bun + Hono instead of Node + Express", "I wanted a compact TypeScript backend with a fast local runtime, simple route composition, middleware, validation integration, tests, and Bun-native WebSocket support. Hono keeps the API small and feature routes are easy to read.", "Node and Express have a larger ecosystem, more examples, and lower hiring-team unfamiliarity. Express would be a conservative choice, but it usually needs more assembly for typing and modern WebSocket handling. The current API is small enough that Hono's minimal surface is useful.", "I did not avoid Express because it is bad; I chose Hono because this API benefits from a smaller typed surface and Bun already provides the runtime and WebSocket server."),
        ("Python worker instead of Go", "The worker uses Ultralytics YOLOv8n and OpenCV. Python has direct, well-supported APIs for both, which minimized model-runtime integration work and matched the fastest path to a working detector.", "Go could produce a smaller operational binary and strong concurrency, but running YOLO would require an external inference runtime, bindings, an ONNX pipeline, or a separate model service. That would move complexity into model conversion and native dependencies.", "I did not use Go because the ML ecosystem was the dominant constraint; Python let me integrate YOLO and OpenCV directly. Go remains a possible control-plane or optimized-worker option later."),
        ("Postgres instead of MongoDB", "Users, cameras, alerts, stats, ownership, and camera deletion have clear relationships. Postgres gives foreign keys, cascading deletes, transactions, JSONB for bounding boxes, timestamp ordering, and indexes for history queries.", "MongoDB offers flexible documents and easy nested data, but this schema is not especially unstructured. I would have to enforce relationships and ownership consistency more in application code. Postgres still accommodates the variable bounding-box object through JSONB.", "I did not choose MongoDB because the core data is relational and benefits from constraints; the one flexible field does not justify giving up those guarantees."),
        ("Redis Streams instead of Kafka or RabbitMQ", "The project needs a lightweight local command/event log between one API and one worker. Redis was already simple to run in Compose, and Streams provide IDs, retention, consumer groups, and acknowledgements.", "Kafka is stronger for high-throughput durable event platforms but adds brokers, partitions, replication, and operational cost far beyond this MVP. RabbitMQ provides mature routing and queues but adds another messaging model and service to learn. Redis Streams are sufficient at current scale, though delivery and trimming need hardening.", "I did not use Kafka because the traffic and team size do not justify Kafka's operational complexity; I needed a small durable stream, not a company-wide event platform."),
        ("MediaMTX + RTSP/WebRTC instead of a simpler stream", "Real cameras commonly expose RTSP, while browsers need a supported low-latency protocol. MediaMTX already implements RTSP ingest, WebRTC serving, WHEP negotiation, and session management.", "MJPEG would be simpler but bandwidth-heavy and inefficient. HLS is widely supported and CDN-friendly but usually adds more latency. Writing WebRTC signaling and media transport myself would be risky and unnecessary. MediaMTX keeps the project focused on product logic.", "I did not use HLS because the demo prioritizes low-latency live monitoring; MediaMTX gives me WebRTC while preserving RTSP compatibility with cameras."),
        ("YOLOv8n instead of a larger or custom model", "The goal is reliable person detection on a local CPU with a quick startup and small download. YOLOv8n is pretrained on COCO, includes the person class, and integrates directly with Ultralytics.", "A larger YOLO model could improve accuracy but would increase CPU latency and memory. A custom model could fit a specific camera environment but needs labeled data, training, evaluation, and model lifecycle work. The event contract lets the model change later without rewriting the app.", "I did not train a custom model because the product question was end-to-end realtime integration; the pretrained person class was enough to validate that system."),
        ("WebSockets instead of polling", "Alerts, FPS, and state changes should appear soon after the worker produces them. A WebSocket lets the API push small events over one long-lived connection.", "Polling is simpler and naturally recovers from failures, but frequent polling wastes requests and adds delay; slow polling makes a surveillance dashboard feel stale. I kept REST for recovery and added WebSocket reconnection with a refresh after reconnecting.", "I did not rely on polling because the dashboard needs low-latency updates, but I still use REST refresh as the recovery path when the socket reconnects."),
        ("Docker Compose instead of Kubernetes", "The project has seven cooperating local services and must be reproducible on one developer machine. Compose provides networking, ports, volumes, environment variables, dependencies, and one-command startup.", "Kubernetes is appropriate for multi-node scheduling, rollouts, service discovery, and self-healing at larger scale. For this local MVP it would add manifests and cluster operations before the product needed them. Optional simple manifests exist, but Kubernetes is not the primary tested runtime.", "I did not lead with Kubernetes because deployment scale did not justify it; Compose solves the actual local orchestration problem with much less operational overhead."),
    ]

    doc.add_heading("Part II: Why This Technology Stack", level=1)
    for idx, (title, chosen, alternative, direct) in enumerate(choices, 1):
        doc.add_heading(f"{idx}. {title}", level=2)
        add_labeled(doc, "Why I chose it", chosen)
        add_labeled(doc, "Realistic alternative and trade-off", alternative)
        add_callout(doc, "One-line answer", direct, color=GOLD)

    doc.add_heading("If I Rebuilt This Today", level=1)
    add_body(doc, "I would keep the broad service boundaries but harden contracts and operations earlier:")
    rebuild = [
        ("Use migrations", "Replace startup CREATE TABLE statements with versioned migrations so schema changes are reviewable and reversible."),
        ("Formalize event schemas", "Validate Redis payloads at both ends and generate shared schemas or contract fixtures instead of relying mainly on documentation."),
        ("Make consumption durable", "Use consumer groups for API event streams, idempotent database constraints, pending-message recovery, dead-letter handling, and stream trimming."),
        ("Design worker ownership", "Add camera leases and unique worker identities before scaling beyond one worker."),
        ("Add evidence snapshots", "Write selected frames to object storage and save the URL already represented by snapshot_url."),
        ("Build alert history UI", "Use the existing cursor/filter API for a searchable history screen with camera and date filters."),
        ("Add observability", "Expose metrics for input FPS, inference latency, reconnect count, stream age, queue lag, and failure rates, with structured logs and tracing IDs."),
        ("Harden internet deployment", "Use managed Postgres/Redis, secret management, strict CORS, TLS, refresh-token strategy, TURN, backups, and network policies."),
        ("Test the complete stack", "Add a Compose integration test that creates a camera, waits for a stored alert, and verifies a browser-consumable stream."),
    ]
    for label, text in rebuild:
        add_labeled(doc, label, text)

    doc.add_heading("Claims To Avoid", level=1)
    add_bullets(doc, [
        "Do not say the worker is Go; it is Python.",
        "Do not say YOLO was trained in this project; a pretrained YOLOv8n model is integrated.",
        "Do not say Kubernetes is the proven deployment platform; Docker Compose is the tested runtime.",
        "Do not claim snapshots, email/SMS alerts, a full alert-history screen, or multi-worker processing.",
        "Do not call development JWT/CORS/WebRTC defaults production-secure.",
    ])
    return save(doc, "02_Challenges_Tradeoffs_and_Tech_Stack.docx")


def build_beginner_guide():
    doc = Document()
    configure_styles(doc)
    add_cover(
        doc,
        "Beginner's Guide",
        "Running, debugging, and understanding the complete local system",
        "GUIDE 03",
        "This guide uses the actual Compose service names. The detection worker is Python; there is no Go worker in the current repository.",
    )
    doc.add_heading("1. The Big Picture", level=1)
    add_body(doc, "This application is not one program. It is seven small programs that cooperate. Docker Compose creates a private network for them, starts them with the right settings, and exposes only the ports your Windows browser or terminal needs.")
    add_code(doc, "Camera/demo -> MediaMTX -> Python worker -> MediaMTX -> browser video\n                              |\n                              +-> Redis -> API -> Postgres + WebSocket -> dashboard")
    add_callout(doc, "Remember", "localhost means your Windows computer. Inside Docker, localhost means that individual container, so containers reach each other using names such as postgres, redis, and mediamtx.")

    doc.add_heading("2. Docker in Plain English", level=1)
    add_labeled(doc, "Docker image", "A packaged recipe containing an application and its dependencies, similar to a reusable machine template.")
    add_labeled(doc, "Container", "A running instance of an image. It is isolated but can use configured ports, files, and networks.")
    add_labeled(doc, "Docker Compose", "The docker-compose.yml file and command that describe and operate all seven containers as one local application.")
    add_labeled(doc, "Why this project uses it", "Installing Postgres, Redis, MediaMTX, FFmpeg, Bun, Node, Python, OpenCV, PyTorch, and YOLO directly on Windows would be fragile. Compose gives each service its correct environment and lets the whole stack start consistently.")

    doc.add_heading("3. What You Need", level=1)
    add_bullets(doc, [
        "Windows with Docker Desktop installed and running.",
        "The project folder at C:\\Users\\vansh\\Desktop\\Smart_camera_surveilance.",
        "Ports 5173, 3000, 5432, 6379, 8554, 8889, 9997, and UDP 8189 available.",
        "Internet access during the first image/model download. Later starts normally use cached images and layers.",
    ])

    doc.add_heading("4. Start the Entire Project", level=1)
    add_numbered(doc, [
        "Open Docker Desktop and wait until Docker Engine is running.",
        "Open PowerShell.",
        "Move into the project directory.",
        "Build changed images and start all normal services in the background.",
        "Check container status, then open the dashboard.",
    ])
    add_code(doc, 'cd "C:\\Users\\vansh\\Desktop\\Smart_camera_surveilance"\ndocker compose up -d --build\ndocker compose ps')
    add_labeled(doc, "Dashboard", "Open http://localhost:5173 and sign in with demo / password.")
    add_labeled(doc, "Normal daily startup", "Use docker compose up -d when code or dependencies have not changed.")
    add_labeled(doc, "Normal shutdown", "Use docker compose down. This removes containers and the Compose network but preserves the named Postgres volume.")
    add_callout(doc, "Data warning", "Do not run docker compose down -v unless you intentionally want to delete stored users, cameras, alerts, and stats.", color=RED)

    doc.add_heading("5. Does Startup Order Matter?", level=1)
    add_body(doc, "You should start the stack with one Compose command. Compose waits for Postgres and Redis health checks before starting the API, and the worker waits for healthy Redis plus a started MediaMTX container. ffmpeg-testcam depends on MediaMTX. A service being started does not always mean it is fully ready, so the worker and browser include retry behavior for temporary video startup delays.")

    doc.add_heading("6. Environment Variables", level=1)
    add_body(doc, "The .env file is optional for local development because docker-compose.yml contains defaults. Copy .env.example to .env only when you want to override them.")
    add_code(doc, "Copy-Item .env.example .env")
    env_rows = [
        ["POSTGRES_USER / PASSWORD / DB", "surveillance", "Local Postgres credentials and database."],
        ["DATABASE_URL", "postgres://...@postgres:5432/...", "API connection string inside Docker."],
        ["REDIS_URL", "redis://redis:6379", "API and worker Redis address."],
        ["JWT_SECRET", "change-me-in-production", "Signs login tokens; replace outside local development."],
        ["API_PORT", "3000", "API listening port."],
        ["PUBLIC_API_URL", "http://localhost:3000", "URL compiled into the frontend for REST."],
        ["PUBLIC_WS_URL", "ws://localhost:3000/ws", "URL compiled into the frontend for events."],
        ["PUBLIC_WEBRTC_BASE_URL", "http://localhost:8889", "Base URL returned for camera WHEP playback."],
        ["ALERT_DEDUPE_SECONDS", "10", "Minimum recent-alert suppression window."],
        ["MEDIAMTX_RTSP_URL", "rtsp://mediamtx:8554", "Worker restream destination."],
        ["WORKER_CONSUMER_NAME", "worker-1", "Redis command consumer identity."],
        ["DETECTION_CONFIDENCE", "0.35", "Minimum YOLO person confidence."],
        ["DETECTION_INTERVAL_MS", "500", "Approximate time between inference runs."],
    ]
    add_table(doc, ["Variable", "Local default", "Purpose"], env_rows, [3000, 3000, 3360])

    doc.add_heading("7. The Seven Services", level=1)
    services = [
        ("frontend", "React/TypeScript UI built by Vite and served by Nginx.", "5173 -> 80"),
        ("api", "Bun/Hono REST and WebSocket server.", "3000"),
        ("worker", "Python/OpenCV/YOLO camera processor.", "No host port"),
        ("postgres", "Durable users, cameras, alerts, and stats.", "5432"),
        ("redis", "Streams for commands, detections, stats, and state.", "6379"),
        ("mediamtx", "RTSP ingest and WebRTC/WHEP video server.", "8554, 8889, 9997, 8189/udp"),
        ("ffmpeg-testcam", "Loops the local person image as an RTSP camera.", "No host port"),
    ]
    add_table(doc, ["Service", "Responsibility", "Host ports"], [list(x) for x in services], [1800, 4800, 2760])

    doc.add_heading("8. Health Checks by Service", level=1)
    add_body(doc, "Start with docker compose ps. Then use the narrowest check for the layer you are investigating.")
    health_rows = [
        ["All", "docker compose ps", "Seven normal services Up; Postgres/Redis healthy."],
        ["frontend", "Invoke-WebRequest http://localhost:5173 -UseBasicParsing", "StatusCode 200."],
        ["api", "Invoke-RestMethod http://localhost:3000/health", "ok = True."],
        ["worker", "docker compose logs --tail=50 worker", "worker ready; camera input connected when active."],
        ["postgres", "docker compose exec postgres pg_isready -U surveillance -d surveillance", "accepting connections."],
        ["redis", "docker compose exec redis redis-cli ping", "PONG."],
        ["mediamtx", "Invoke-RestMethod http://localhost:9997/v3/paths/list", "JSON path list; testcam when published."],
        ["ffmpeg-testcam", "docker compose logs --tail=30 ffmpeg-testcam", "Output to rtsp://mediamtx:8554/testcam; no Conversion failed."],
    ]
    add_table(doc, ["Layer", "Command", "Healthy signal"], health_rows, [1500, 4920, 2940])

    doc.add_heading("9. Viewing Logs", level=1)
    add_body(doc, "Logs are the main way to observe containers that do not have a web page. -f follows new lines until you press Ctrl+C. --tail limits old output, and --since focuses on a recent incident.")
    add_code(doc, "docker compose logs -f api\ndocker compose logs -f worker\ndocker compose logs -f postgres\ndocker compose logs -f redis\ndocker compose logs -f mediamtx\ndocker compose logs -f ffmpeg-testcam\ndocker compose logs -f frontend")
    add_code(doc, "docker compose logs --tail=100 worker\ndocker compose logs --since=5m mediamtx")

    doc.add_heading("10. Test the Main Functionality", level=1)
    add_numbered(doc, [
        "Open the dashboard and log in.",
        "Confirm the test camera exists or add rtsp://mediamtx:8554/testcam.",
        "Click Start and allow YOLO to load on the first run.",
        "Confirm the person image, roughly 25 FPS, nonzero detections, and recent alerts.",
        "Open Stream diagnostics and expect Requested Live, Worker Live, Browser Playing, Input Enabled.",
        "Leave the page open and confirm timestamps update without Refresh.",
        "Stop and start the camera and observe Connecting before Playing.",
        "Refresh or log out/in and confirm cameras and stored recent alerts remain.",
    ])

    doc.add_heading("11. Run Automated Tests", level=1)
    add_code(doc, 'docker compose build api frontend worker\ndocker compose run --rm --no-deps api sh -c "bun run typecheck && bun test"\ndocker compose run --rm --build frontend-test\ndocker compose run --rm --no-deps worker python -m pytest -q')
    add_body(doc, "The API tests use mocked database and Redis boundaries for route/ownership behavior, so they do not modify your real cameras. Manual testing is still required for real video, network, and inference behavior.")

    doc.add_heading("12. A Beginner's Debugging Method", level=1)
    add_numbered(doc, [
        "Read the visible state: Requested, Worker, Browser, FPS, and Live updates.",
        "Check docker compose ps before assuming an application bug.",
        "Test the API health URL separately from video.",
        "Follow the video pipeline from source to browser and stop at the first failing layer.",
        "Read fresh logs with --since or --tail instead of scanning hours of output.",
        "Change one layer at a time, retest, and keep the evidence.",
    ])
    add_code(doc, "source camera -> MediaMTX testcam -> worker input -> camera-specific MediaMTX path -> WebRTC browser")

    errors = [
        ("Port is already allocated", "Docker cannot bind a host port because another process or container uses it.", "Run docker compose ps and Get-NetTCPConnection -LocalPort <port>. Stop the conflicting application or change the host-side port in Compose."),
        ("Docker daemon is not running", "Commands cannot contact Docker Desktop.", "Open Docker Desktop, wait for the engine, then retry docker compose up -d."),
        ("Container exits or restarts", "The service process crashed and restart policy may be looping it.", "Run docker compose ps and docker compose logs --tail=100 <service>. Fix the first error, then recreate only that service."),
        ("Database connection refused", "The API cannot reach Postgres or Postgres is not ready.", "Confirm postgres is healthy, run pg_isready inside it, and verify DATABASE_URL uses host postgres rather than localhost inside Docker."),
        ("Redis connection failure", "Commands and events cannot move between API and worker.", "Run redis-cli ping, inspect API/worker logs, and confirm REDIS_URL is redis://redis:6379."),
        ("Failed to resolve hostname mediamtx", "Docker's internal DNS/network cannot resolve the service name.", "Check both containers are in the same Compose project. Recreate the network with docker compose down followed by docker compose up -d. Do not use -v."),
        ("Worker stays Connecting", "The start command was accepted, but the worker has not opened frames or published live state.", "Check worker and ffmpeg-testcam logs. Run ffprobe from the worker against testcam. Redis XINFO GROUPS camera.commands distinguishes command lag from stream opening."),
        ("WebRTC negotiation failed", "The browser reached MediaMTX before the camera-specific output existed, or port/ICE configuration is wrong.", "Allow the built-in retries. If it lasts over ten seconds, inspect MediaMTX for no one is publishing, confirm 8889/tcp and 8189/udp, and compare Worker versus Browser diagnostics."),
        ("Live video but stale FPS/alerts", "The media path works independently, but WebSocket events are disconnected.", "Check the Live updates indicator. The frontend should reconnect and refresh. Inspect API /ws logs and PUBLIC_WS_URL if it does not."),
        ("YOLO confidence barely changes", "The bundled demo repeats the exact same still image.", "This is expected. Replace the demo input with a moving MP4 or a real RTSP camera to vary pose, lighting, scale, and confidence."),
        ("Tests cannot import worker modules", "Pytest was started without /app on Python's import path or an old image lacks pytest.ini.", "Rebuild worker and run docker compose run --rm --no-deps worker python -m pytest -q."),
    ]
    doc.add_heading("13. Common Errors and What To Do", level=1)
    for title, cause, fix in errors:
        doc.add_heading(title, level=2)
        add_labeled(doc, "Meaning", cause)
        add_labeled(doc, "Diagnosis and fix", fix)

    doc.add_heading("14. How to Test Another Video", level=1)
    add_numbered(doc, [
        "Place a file named test-video.mp4 in demo-assets.",
        "In ffmpeg-testcam, replace the still-image input arguments with -stream_loop, -1, -re, -i, /demo-assets/test-video.mp4.",
        "Recreate only the source with docker compose up -d --force-recreate ffmpeg-testcam.",
        "Keep the dashboard camera URL as rtsp://mediamtx:8554/testcam, then Stop and Start it.",
    ])
    add_body(doc, "A real IP camera can be registered with its own RTSP URL, but the Docker host must be able to reach that camera's network and credentials must be handled carefully.")

    doc.add_heading("15. Future Scope", level=1)
    future = [
        ("Alert history UI", "Use the existing camera/date/cursor filters from GET /alerts to build a searchable page and Load more flow."),
        ("Snapshots", "Save selected detection frames to S3-compatible object storage and persist the URL in the existing snapshot_url column."),
        ("Notifications", "Consume stored alert events in a notification worker and send email, push, SMS, or webhooks with per-user preferences and rate limits."),
        ("Multiple detection workers", "Add camera leases or consistent assignment, unique worker identities, pending-command recovery, and capacity-aware scheduling."),
        ("Model optimization", "Export to ONNX/TensorRT or use GPU inference, measure latency, and batch only when the latency trade-off is acceptable."),
        ("Per-camera rules", "Store confidence, detection classes, schedules, and regions of interest per camera and include them in Start commands."),
        ("Cloud deployment", "Use managed Postgres/Redis, object storage, TLS ingress, secret management, backups, TURN, and monitoring."),
        ("Multi-tenant administration", "Add organizations, roles, camera sharing, audit logs, refresh tokens, account recovery, and stricter policies."),
        ("Observability", "Publish metrics for camera age, FPS, inference latency, reconnects, Redis lag, alert rate, and errors; add structured logs and tracing."),
        ("End-to-end CI", "Start Compose in CI, create a test camera, wait for an alert, verify Postgres and WebSocket behavior, and tear down without deleting developer data."),
    ]
    for label, text in future:
        add_labeled(doc, label, text)

    doc.add_heading("16. Vocabulary Cheat Sheet", level=1)
    vocab = [
        ["RTSP", "A camera-oriented protocol used to request/control a media stream."],
        ["WebRTC", "Browser-supported low-latency realtime media transport."],
        ["WHEP", "An HTTP-based way for a browser to receive a WebRTC stream."],
        ["REST", "Request-response API calls for actions and data."],
        ["WebSocket", "A persistent two-way connection used here for pushed alerts and stats."],
        ["JWT", "A signed token proving the logged-in user's identity until expiry."],
        ["Redis Stream", "An ordered retained sequence of message entries with IDs."],
        ["Consumer group", "A Redis mechanism for distributing and acknowledging stream work."],
        ["YOLOv8n", "A small pretrained object detector; this project filters it to person."],
        ["Inference", "Running an already-trained model on a frame to get predictions."],
        ["Container", "An isolated running application created from an image."],
        ["Volume", "Docker-managed persistent storage; Postgres data survives container recreation."],
    ]
    add_table(doc, ["Term", "Plain-English meaning"], vocab, [2100, 7260])
    return save(doc, "03_Beginners_Guide_Running_Debugging_and_Concepts.docx")


if __name__ == "__main__":
    paths = [build_questions(), build_tradeoffs(), build_beginner_guide()]
    for path in paths:
        print(path)
